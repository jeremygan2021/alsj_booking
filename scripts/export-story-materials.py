#!/usr/bin/env python3
"""Export private actor packets and commissioning briefs from the runtime story source."""
from pathlib import Path
import json, html, shutil, csv
ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'server/story-materials.json').read_text())
OUT=ROOT/'交付包_阿拉斯加之夜'
NPC=OUT/'NPC独立物料'; MODEL=OUT/'3D模型制作需求'
NPC.mkdir(parents=True,exist_ok=True);MODEL.mkdir(parents=True,exist_ok=True)
def escaped(x):return html.escape(str(x))
def paragraph(x):return '<p>'+escaped(x).replace('\n','<br>')+'</p>'
css='''body{font:16px/1.8 -apple-system,BlinkMacSystemFont,"PingFang SC",sans-serif;color:#172b3d;background:#f2f5f7;margin:0}main{max-width:900px;margin:24px auto;padding:32px;background:white}header{display:flex;gap:24px;align-items:center;border-bottom:2px solid #1d425b;padding-bottom:24px}img{width:150px;height:150px;object-fit:cover;border-radius:8px}h1{font-size:28px;margin:0}h2{font-size:20px;color:#1d425b;margin-top:32px}small{color:#566575}table{width:100%;border-collapse:collapse}td,th{border:1px solid #ccd4dc;padding:8px;text-align:left;vertical-align:top}button{padding:10px 18px;background:#1d425b;color:white;border:0;border-radius:5px}a{color:#1d425b}@page{size:A4;margin:16mm}@media print{body{background:white;font-size:11pt}main{padding:0;margin:0;max-width:none}button{display:none}header{break-inside:avoid}h2{break-after:avoid}tr{break-inside:avoid}img{width:100px;height:100px}}'''
index=[]
for n in DATA['npcs']:
 folder=NPC/f"{n['npcId']}_{n['name']}";folder.mkdir(exist_ok=True)
 shutil.copy2(ROOT/n['portrait'],folder/'头像.webp')
 supp=[s for s in n['sections'] if s['title'] not in ('公开人物小传','第一人称故事卡')]
 md=f"# {n['npcId']}｜{n['name']}\n\n{n['english']}｜{n['age']}岁｜{n['role']}\n\n**演员本人/GM保密物料**\n\n![代码现有头像](头像.webp)\n\n## 公开故事\n\n{n['publicBio']}\n\n## 第一人称故事\n\n{n['story']}\n\n## 我知道的内容\n\n"+'\n'.join('- '+x for x in n['knowledge'])+'\n\n## 我不知道与不能说的内容\n\n'+'\n'.join('- '+x for x in n['unknown'])+'\n\n## 我的时间线和走位\n\n| 时间 | 事件 |\n|---|---|\n'+'\n'.join(f"| {x['time']} | {x['event']} |" for x in n['timeline'])+'\n\n'
 for s in supp:md+=f"## {s['title']}\n\n{s['text']}\n\n"
 md+='## 证据编号对照\n\n| 原稿编号 | 代码编号 | 名称 |\n|---|---|---|\n'
 relevant=[e for e in DATA['evidence'] if e['id'] in n['evidenceIds']]
 for e in relevant:md+=f"| {e['id']} | {e['clueId']} | {e['title']} |\n"
 md+=f"\n## 头像说明\n\n{n['portraitNote']} 原代码资源：`{n['portrait']}`。\n"
 with (folder/'角色道具清单.csv').open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f);w.writerow(['NPC编号','NPC','道具','制作状态','头像与剧本来源'])
  props=next((s['text'] for s in n['sections'] if s['title']=='专属道具'),'')
  for line in props.splitlines():
   if line.startswith('- '):w.writerow([n['npcId'],n['name'],line[2:],'待制作','角色原稿及代码现有头像'])
 (folder/'演员剧本.md').write_text(md)
 content=f"<header><img src='头像.webp' alt='{escaped(n['name'])}头像'><div><small>{n['npcId']} · 演员保密</small><h1>{escaped(n['name'])}</h1><p>{escaped(n['english'])} · {n['age']}岁<br>{escaped(n['role'])}</p></div></header><button onclick='window.print()'>打印角色物料</button>"
 for title,text in [('公开故事',n['publicBio']),('第一人称故事',n['story'])]:content+=f'<h2>{title}</h2>'+paragraph(text)
 for title,items in [('我知道的内容',n['knowledge']),('我不知道与不能说的内容',n['unknown'])]:content+=f'<h2>{title}</h2><ul>'+''.join('<li>'+escaped(x)+'</li>' for x in items)+'</ul>'
 content+='<h2>我的时间线与走位</h2><table><tr><th>时间</th><th>事件</th></tr>'+''.join(f"<tr><td>{escaped(t['time'])}</td><td>{escaped(t['event'])}</td></tr>" for t in n['timeline'])+'</table>'
 for s in supp:content+='<h2>'+escaped(s['title'])+'</h2>'+paragraph(s['text'])
 content+='<h2>证据编号对照</h2><table><tr><th>原稿</th><th>代码</th><th>名称</th></tr>'+''.join(f"<tr><td>{e['id']}</td><td>{e['clueId']}</td><td>{escaped(e['title'])}</td></tr>" for e in relevant)+'</table>'
 content+='<h2>头像说明</h2>'+paragraph(n['portraitNote'])
 (folder/'可打印演员物料.html').write_text(f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escaped(n["name"])}演员物料</title><style>{css}</style><main>{content}</main></html>')
 index.append(f"<li><a href='{folder.name}/可打印演员物料.html'>{n['npcId']} · {escaped(n['name'])}</a> · {escaped(n['role'])}</li>")
(NPC/'index.html').write_text(f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>NPC物料索引</title><style>{css}</style><main><h1>六名NPC独立物料</h1><p>含现有代码头像、故事、时间线、知道与不知道的内容、分阶段证言、任务和道具。只将对应角色文件夹交给本人演员；完整目录由GM保存。</p><ul>{"".join(index)}</ul></main></html>')
with (MODEL/'证据模型映射.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f);w.writerow(['原稿ID','代码ID','证据','场景','模型ID','正式GLB','阶段','关键证据'])
 models={m['id']:m for m in DATA['models']}
 for e in DATA['evidence']:w.writerow([e['id'],e['clueId'],e['title'],e['area'],e['modelId'],models[e['modelId']]['filename'],e['phase'],e['critical']])
modeldoc='# 线索3D模型制作需求\n\n13类正式程序化GLB已生成，支撑25条原稿证据与6个AR调查点。代码保留现有3类示例GLB作为缺失时的回退。\n\n原稿EV编号与代码E编号已显式映射，不能直接按数字对应。E-09为保留的拍卖手枪排除线索，原稿证据映射到其余25个编号。\n\n'
for m in DATA['models']:
 text=f"## {m['name']}\n\n- 模型ID：{m['id']}\n- 状态：{m['status']}\n- 正式文件：`public/models/{m['filename']}`\n- 对应原稿证据：{'、'.join(m['evidenceIds']) or 'AR排除性调查'}\n- 对应代码线索：{'、'.join(m['clueIds']) or '待配置独立AR节点'}\n- 场景：{m['scene']}\n\n### 外形与可观察细节\n\n{m['geometry']}\n\n### 交互热点\n\n{'、'.join(m['hotspots'])}\n\n{m['interaction']}\n\n### 信息与剧情边界\n\n{m['limits']}\n\n### 格式与交付\n\n{m['format']}\n\n### 可直接交给3D生成工具的提示词\n\n{m['prompt']}\n\n"
 (MODEL/(m['id']+'.md')).write_text('# '+m['name']+'制作单\n\n'+text)
 modeldoc+=text
modeldoc+='''## 代码接入流程\n\n1. 对照清单制作并审查GLB，确保基础网格不含未解锁证据正文。\n2. 将GLB按固定文件名放到`public/models/`；服务器固定清单检测文件存在后自动替换示例模型。\n3. 查看器仍沿用现有二维码/证据入口、鉴定权限；不会因为换模型跳过证据资格检查。\n4. 本次已接入正式模型、观察点标记、镜头定位与自动回退；动态模型状态及分层报告叠加仍可后续扩展。\n5. AR-01至AR-06均有公开调查节点入口，使用/ar?node=AR-01等链接；墙面图像追踪场景仍由现有AR编辑器配置。\n'''
(MODEL/'3D模型制作总需求.md').write_text(modeldoc)
(OUT/'NPC与3D物料使用说明.md').write_text('''# NPC与3D物料使用说明\n\n- 从`NPC独立物料/index.html`打开六名角色包，每人独立文件夹含代码头像、Markdown剧本和可打印HTML。\n- NPC后台增加“我的NPC剧本”；GM“NPC档案”里查看演员剧本和模型制作要求。演员只能收到自己剧本，GM可查看全部。\n- 从`3D模型制作需求/3D模型制作总需求.md`查看13类模型要求；每类另有独立制作单及完整证据编号CSV。\n- 13类正式GLB已生成；已接入观察点定位、固定文件名与示例模型回退。\n- 当前使用源为`server/story-materials.json`；更新后运行`npm run export:materials`重新导出物料。\n- 本次以用户提供故事稿为准统一角色公开身份，保留代码旧ID，头像全部沿用代码现有文件，不重绘。\n- 新版本迁移仅更新未被GM改动的旧默认文字；已有自定义文案需GM对照新版资料审核。\n''')
# Update the previously delivered workbook with actor-specific and model authoring tables.
try:
 from openpyxl import load_workbook,Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 p=OUT/'阿拉斯加之夜_制作管理表.xlsx'
 wb=load_workbook(p) if p.exists() else Workbook()
 for name in ['NPC演员索引','NPC个人时间线','NPC已知与边界','3D模型制作要求','证据代码映射']:
  if name in wb:del wb[name]
 def sheet(name,headers,rows,widths):
  s=wb.create_sheet(name);s.append(headers)
  for row in rows:s.append(row)
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
  for i,width in enumerate(widths,1):s.column_dimensions[get_column_letter(i)].width=width
  for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
  for row in s:
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
  return s
 sheet('NPC演员索引',['原稿编号','代码ID','姓名','身份','头像代码位置','人物公开故事','演员秘密'],[[n['npcId'],n['id'],n['name'],n['role'],n['portrait'],n['publicBio'],n['secret']] for n in DATA['npcs']],[14,16,20,32,38,70,65])
 sheet('NPC个人时间线',['NPC','时间','本人事件/演出走位'],[[n['name'],t['time'],t['event']] for n in DATA['npcs'] for t in n['timeline']],[20,22,90])
 sheet('NPC已知与边界',['NPC','类型','内容'],[[n['name'],kind,text] for n in DATA['npcs'] for kind,values in [('已知事实',n['knowledge']),('不知道/不能说',n['unknown'])] for text in values],[20,24,95])
 sheet('3D模型制作要求',['模型ID','模型名','文件名','原稿证据','代码线索','场景','外形','热点','交互要求','剧情边界','格式','生成提示词','状态'],[[m['id'],m['name'],m['filename'],'、'.join(m['evidenceIds']),'、'.join(m['clueIds']),m['scene'],m['geometry'],'、'.join(m['hotspots']),m['interaction'],m['limits'],m['format'],m['prompt'],m['status']] for m in DATA['models']],[24,30,32,35,35,26,65,55,65,65,65,90,14])
 sheet('证据代码映射',['原稿ID','代码ID','名称','基础事实','进阶提示','阶段','对应模型'],[[e['id'],e['clueId'],e['title'],e['body'],e['detail'],e['phase'],e['modelId']] for e in DATA['evidence']],[14,14,32,80,80,14,30])
 # Repair earlier AR table's sixth column: older export omitted this header.
 if 'AR节点' in wb:
  s=wb['AR节点']
  if s.max_column==6 and s.cell(1,6).value is None:
   s.cell(1,6,'剧情/制作边界')
   s.column_dimensions['F'].width=55
   for t in s.tables.values():
    if len(t.tableColumns)==5:
     from openpyxl.worksheet.table import TableColumn
     t.tableColumns.append(TableColumn(id=6,name='剧情/制作边界'))
 wb.save(p)
except ImportError:
 print('openpyxl unavailable; HTML/Markdown/CSV exported, workbook not updated.')
print('Exported 6 portrait actor packets and 13 independent model briefs to',OUT)

# Each actor can receive only their own archive; the complete archive is for GM/production.
from zipfile import ZipFile, ZIP_DEFLATED
for folder in sorted(NPC.glob('N*_*')):
 if not folder.is_dir():continue
 with ZipFile(NPC/(folder.name+'_演员物料.zip'),'w',ZIP_DEFLATED) as z:
  for file in sorted(folder.iterdir()):
   if file.is_file():z.write(file,file.relative_to(NPC))
with ZipFile(OUT/'NPC与3D制作物料包.zip','w',ZIP_DEFLATED) as z:
 for directory in [NPC,MODEL]:
  for file in sorted(directory.rglob('*')):
   if file.is_file() and file.suffix!='.zip':z.write(file,file.relative_to(OUT))
 for file in [OUT/'阿拉斯加之夜_制作管理表.xlsx',OUT/'NPC与3D物料使用说明.md']:
  if file.exists():z.write(file,file.name)
